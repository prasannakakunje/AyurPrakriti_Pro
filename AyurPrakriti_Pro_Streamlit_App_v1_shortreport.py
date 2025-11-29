# AyurPrakriti_Pro_Streamlit_App.py
# Full Streamlit app — AyurPrakriti Pro (Prakriti/Vikriti + Psychometrics)
# Key fixes: DejaVu font, scaled logo, unique ReportLab styles, robust fallback PDF

import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from io import BytesIO
from datetime import datetime, timedelta
import os, json, shutil, logging, traceback
import sqlite3
from pathlib import Path
from passlib.context import CryptContext
import yaml
from docx import Document

# ReportLab imports
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Image as RLImage,
    Table,
    TableStyle,
    PageBreak,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

# ---------------- APP DIRECTORIES & BRANDING ------------------------------
APP_DIR = Path.home() / ".ayurprakriti_app"
APP_DIR.mkdir(parents=True, exist_ok=True)
FONTS_DIR = APP_DIR / "fonts"
FONTS_DIR.mkdir(parents=True, exist_ok=True)
TMP_DIR = APP_DIR / "tmp"
TMP_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = APP_DIR / "ayurprakriti.db"
CFG_PATH = APP_DIR / "config_rules.yaml"

# Pre-copy a dev logo if present in container path
if Path("/mnt/data/logo.png").exists():
    try:
        shutil.copy("/mnt/data/logo.png", APP_DIR / "logo.png")
    except Exception:
        pass

BRAND = {
    "clinic_name": "Kakunje Wellness",
    "tagline": "Authentic Ayurveda | Modern Precision",
    "doctor": "Prof. Dr. Prasanna Kakunje, MD (Ayu), (PhD)",
    "address": "Janani Complex, Nagarakatte Road, Moodbidri, Karnataka",
    "phone": "+91-9483697676",
    "email": "prasanna@kakunje.com",
    "website": "https://kakunje.com",
    "accent_color": "#0F7A61",
}

# ---------------- LOGGER ---------------------------------------------------
logger = logging.getLogger("ayurprakriti_app")
if not logger.handlers:
    fh = logging.FileHandler(APP_DIR / "app_debug.log")
    fh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger.addHandler(fh)
logger.setLevel(logging.INFO)

# ---------------- CONFIG & DEFAULTS ---------------------------------------
pwd_context = CryptContext(schemes=["pbkdf2_sha256"], deprecated="auto")
DEFAULT_CFG = {
    "meta": {
        "app_name": "AyurPrakriti Pro",
        "version": "1.0",
        "author": "Auto-generated for Dr. Prasanna",
    },
    "questions": {
        "prakriti": [
            {
                "id": "P1",
                "text": "Natural body frame: thin/slender",
                "weights": {"Vata": 1.0},
            },
            {
                "id": "P2",
                "text": "Tendency for dry, rough skin",
                "weights": {"Vata": 1.0},
            },
            {
                "id": "P3",
                "text": "Tendency for warm, reddish skin",
                "weights": {"Pitta": 1.0},
            },
            {
                "id": "P4",
                "text": "Heavier, solid body build",
                "weights": {"Kapha": 1.0},
            },
            {
                "id": "P5",
                "text": "Sleep depth & continuity (deep = higher)",
                "weights": {"Kapha": 1.0},
            },
            {
                "id": "P6",
                "text": "Digestion: variable/irregular vs steady",
                "weights": {"Vata": 1.0},
            },
            {
                "id": "P7",
                "text": "Perspiration: tends to sweat easily",
                "weights": {"Pitta": 1.0},
            },
            {
                "id": "P8",
                "text": "Memory: quick/active memory",
                "weights": {"Vata": 0.6, "Pitta": 0.4},
            },
            {
                "id": "P9",
                "text": "Body temperature: preference for cool climates",
                "weights": {"Pitta": 0.8},
            },
            {
                "id": "P10",
                "text": "Build: tendency to gain weight easily",
                "weights": {"Kapha": 1.0},
            },
        ],
        "vikriti": [
            {
                "id": "V1",
                "text": "Anxiety, restlessness today",
                "weights": {"Vata": 1.0},
            },
            {
                "id": "V2",
                "text": "Anger, irritability, impatience today",
                "weights": {"Pitta": 1.0},
            },
            {
                "id": "V3",
                "text": "Heaviness, lethargy, congestion today",
                "weights": {"Kapha": 1.0},
            },
            {
                "id": "V4",
                "text": "Loose stools or irregular digestion",
                "weights": {"Vata": 1.0},
            },
            {
                "id": "V5",
                "text": "Acidity, excess heat, heartburn",
                "weights": {"Pitta": 1.0},
            },
            {
                "id": "V6",
                "text": "Nasal congestion, mucus, slow clearance",
                "weights": {"Kapha": 1.0},
            },
            {
                "id": "V7",
                "text": "Insomnia or fragmented sleep tonight",
                "weights": {"Vata": 0.8},
            },
            {
                "id": "V8",
                "text": "Excess thirst or dry mouth",
                "weights": {"Pitta": 0.6},
            },
        ],
        "psychometric": [
            {"id": "E1", "text": "Extraverted, enthusiastic"},
            {"id": "E6", "text": "Reserved, quiet"},
            {"id": "A1", "text": "Critical, quarrelsome"},
            {"id": "A6", "text": "Sympathetic, warm"},
            {"id": "C1", "text": "Dependable, self-disciplined"},
            {"id": "C6", "text": "Disorganized, careless"},
            {"id": "N1", "text": "Anxious, easily upset"},
            {"id": "N6", "text": "Calm, emotionally stable"},
            {"id": "O1", "text": "Open to new experiences, complex"},
            {"id": "O6", "text": "Conventional, uncreative"},
        ],
    },
    "mappings": {
        "career_rules": {
            "Vata": ["Writer, Designer, Creative Entrepreneur, Researcher"],
            "Pitta": ["Clinician, Analyst, Manager, Engineer, Competitive roles"],
            "Kapha": ["Teacher, Counselor, Hospitality, Agriculture, HR"],
        },
        "dosha_thresholds": {"mild": 55, "moderate": 70, "severe": 85},
    },
}

if not CFG_PATH.exists():
    with open(CFG_PATH, "w", encoding="utf-8") as f:
        yaml.safe_dump(DEFAULT_CFG, f, sort_keys=False)
with open(CFG_PATH, "r", encoding="utf-8") as f:
    CONFIG = yaml.safe_load(f)

# ---------------- DATABASE -------------------------------------------------
conn = sqlite3.connect(str(DB_PATH), check_same_thread=False)
cur = conn.cursor()
cur.executescript(
    """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE,
    display_name TEXT,
    password_hash TEXT,
    role TEXT DEFAULT 'clinician',
    created_at TEXT
);
CREATE TABLE IF NOT EXISTS patients (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT,
    age INTEGER,
    gender TEXT,
    contact TEXT,
    created_at TEXT
);
CREATE TABLE IF NOT EXISTS assessments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    patient_id INTEGER,
    assessor TEXT,
    data_json TEXT,
    created_at TEXT,
    FOREIGN KEY(patient_id) REFERENCES patients(id)
);
"""
)
conn.commit()

cur.execute("SELECT COUNT(1) FROM users")
if cur.fetchone()[0] == 0:
    ph = pwd_context.hash("admin123")
    cur.execute(
        "INSERT INTO users (username, display_name, password_hash, role, created_at) VALUES (?,?,?,?,?)",
        ("admin", "Administrator", ph, "admin", datetime.now().isoformat()),
    )
    conn.commit()


# ---------------- HELPERS: DB & SCORING ------------------------------------
def verify_user(username, password):
    cur.execute(
        "SELECT password_hash, display_name, role FROM users WHERE username=?",
        (username,),
    )
    r = cur.fetchone()
    if not r:
        return False, None
    ph, display, role = r
    return (pwd_context.verify(password, ph), {"display_name": display, "role": role})


def create_patient(name, age, gender, contact):
    cur.execute(
        "INSERT INTO patients (name, age, gender, contact, created_at) VALUES (?,?,?,?,?)",
        (name, age, gender, contact, datetime.now().isoformat()),
    )
    conn.commit()
    return cur.lastrowid


def save_assessment(patient_id, assessor, data):
    cur.execute(
        "INSERT INTO assessments (patient_id, assessor, data_json, created_at) VALUES (?,?,?,?)",
        (
            patient_id,
            assessor,
            json.dumps(data, ensure_ascii=False),
            datetime.now().isoformat(),
        ),
    )
    conn.commit()
    return cur.lastrowid


def load_patients():
    return pd.read_sql_query("SELECT * FROM patients ORDER BY created_at DESC", conn)


def load_assessments(patient_id=None):
    if patient_id:
        return pd.read_sql_query(
            "SELECT * FROM assessments WHERE patient_id=? ORDER BY created_at DESC",
            conn,
            params=(patient_id,),
        )
    return pd.read_sql_query("SELECT * FROM assessments ORDER BY created_at DESC", conn)


def score_dosha_from_answers(answers, question_list):
    totals = {"Vata": 0.0, "Pitta": 0.0, "Kapha": 0.0}
    for q in question_list:
        qid = q["id"]
        w = q.get("weights", {})
        val = answers.get(qid, 3)
        for d in totals:
            totals[d] += w.get(d, 0) * float(val)
    s = sum(totals.values())
    if s <= 0:
        return {k: round(100 / 3, 1) for k in totals}
    pct = {k: round((v / s) * 100, 1) for k, v in totals.items()}
    return pct


def psychometric_tipiscale(answers):
    try:
        ext = (answers["E1"] + (8 - answers["E6"])) / 2.0
        agr = (((8 - answers["A1"]) + answers["A6"])) / 2.0
        con = (answers["C1"] + (8 - answers["C6"])) / 2.0
        emo = (answers["N1"] + (8 - answers["N6"])) / 2.0
        ope = (answers["O1"] + (8 - answers["O6"])) / 2.0
    except Exception:
        return {
            "Extraversion": 50,
            "Agreeableness": 50,
            "Conscientiousness": 50,
            "Emotionality": 50,
            "Openness": 50,
        }
    raw = {
        "Extraversion": ext,
        "Agreeableness": agr,
        "Conscientiousness": con,
        "Emotionality": emo,
        "Openness": ope,
    }
    return {k: round((v - 1) / 6 * 100, 1) for k, v in raw.items()}


def recommend_career(dosha_percent, psycho_pct, cfg=CONFIG):
    dom = max(dosha_percent, key=dosha_percent.get)
    base = cfg["mappings"]["career_rules"].get(dom, [])
    recs = []
    for r in base:
        score = 50
        if psycho_pct.get("Conscientiousness", 50) > 65 and "Manager" in r:
            score += 10
        if psycho_pct.get("Openness", 50) > 60 and ("Research" in r or "Creative" in r):
            score += 10
        if psycho_pct.get("Extraversion", 50) > 60 and (
            "Sales" in r or "Clinician" in r or "Manager" in r
        ):
            score += 8
        recs.append(
            {
                "role": r,
                "score": score,
                "reason": f"Matches dominant {dom} + psychometric signals.",
            }
        )
    if psycho_pct.get("Openness", 50) > 70 and "Research & Innovation" not in [
        x["role"] for x in recs
    ]:
        recs.append(
            {
                "role": "Research & Innovation / R&D",
                "score": 65,
                "reason": "High openness suggests research fit.",
            }
        )
    return sorted(recs, key=lambda x: -x["score"])


def recommend_relationship(dosha_pct, psycho_pct):
    tips = []
    dom = max(dosha_pct, key=dosha_pct.get)
    if dom == "Vata":
        tips.append(
            (
                "Stability & routine",
                "Vata benefits from grounding routines and predictable schedules.",
            )
        )
    if dom == "Pitta":
        tips.append(
            (
                "Cooling communication",
                "Pause before responding; use neutral language; avoid public challenges.",
            )
        )
    if dom == "Kapha":
        tips.append(
            (
                "Stimulate & vary",
                "Introduce gentle novelty and clear plans to reduce inertia.",
            )
        )
    if psycho_pct.get("Agreeableness", 50) < 40:
        tips.append(
            (
                "Reflective listening",
                "Practice summarizing partner words before replying.",
            )
        )
    if psycho_pct.get("Emotionality", 50) > 60:
        tips.append(
            (
                "Emotion regulation",
                "Simple breathing exercises and journaling help before tough talks.",
            )
        )
    return tips


def recommend_health(dosha_pct, vikriti_pct, cfg=CONFIG):
    dom = max(dosha_pct, key=dosha_pct.get)
    combined = {
        d: round((dosha_pct[d] + vikriti_pct.get(d, 0)) / 2, 1) for d in dosha_pct
    }
    thresholds = cfg["mappings"]["dosha_thresholds"]
    severity = {}
    for d, v in combined.items():
        if v >= thresholds["severe"]:
            severity[d] = "severe"
        elif v >= thresholds["moderate"]:
            severity[d] = "moderate"
        elif v >= thresholds["mild"]:
            severity[d] = "mild"
        else:
            severity[d] = "balanced"
    rec = {"diet": [], "lifestyle": [], "herbs": [], "severity": severity}
    if dom == "Vata":
        rec["diet"] = [
            "Warm, cooked meals; include healthy oils; regular mealtimes; avoid cold/raw foods"
        ]
        rec["lifestyle"] = [
            "Daily Abhyanga (oil massage), grounding routines, regular sleep (approx 10pm-6am), gentle yoga"
        ]
        rec["herbs"] = [
            "Ashwagandha — consult physician for dosing",
            "Bala for strength — clinical supervision",
        ]
    if dom == "Pitta":
        rec["diet"] = [
            "Cooling foods; reduce spicy & fried items; include bitter greens & cooling herbs"
        ]
        rec["lifestyle"] = [
            "Cooling pranayama, avoid midday heat exposure, calming walks"
        ]
        rec["herbs"] = ["Amla, Guduchi — discuss with physician"]
    if dom == "Kapha":
        rec["diet"] = [
            "Light, dry foods; reduce dairy & sweets; favor warm spiced items like trikatu"
        ]
        rec["lifestyle"] = [
            "Stimulating exercise 30-60 min daily, dry massage (udvartana), varied schedules"
        ]
        rec["herbs"] = ["Trikatu, Guggulu — clinical supervision required"]
    return rec


# ---------------- CHART HELPER ---------------------------------------------
def _make_bar_chart(series: dict, title: str, filename: Path):
    plt.close("all")
    keys = list(series.keys())
    vals = [series[k] for k in keys]
    fig, ax = plt.subplots(figsize=(6, 2.4))
    palette = ["#6fbf73", "#f5a623", "#6fb0d9"]
    bars = ax.bar(keys, vals, color=palette[: len(keys)])
    ax.set_ylim(0, 100)
    ax.set_ylabel("Percent")
    ax.set_title(title)
    for bar, v in zip(bars, vals):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            v + 1,
            f"{v}%",
            ha="center",
            va="bottom",
            fontsize=8,
        )
    fig.tight_layout()
    fig.savefig(filename, dpi=150)
    plt.close(fig)


# ---------------- FONT REGISTRATION ----------------------------------------
DEJAVU_PATH = None
_fonts = list(FONTS_DIR.glob("DejaVuSans*.ttf"))
if _fonts:
    exact = FONTS_DIR / "DejaVuSans.ttf"
    if exact.exists():
        DEJAVU_PATH = str(exact)
    else:
        DEJAVU_PATH = str(_fonts[0])
else:
    # typical system paths
    for cand in [
        r"C:\Windows\Fonts\DejaVuSans.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/DejaVuSans.ttf",
    ]:
        if os.path.exists(cand):
            DEJAVU_PATH = cand
            break

if DEJAVU_PATH:
    try:
        pdfmetrics.registerFont(TTFont("DejaVuSans", DEJAVU_PATH))
        # try register bold if available
        bold_path = Path(DEJAVU_PATH).with_name("DejaVuSans-Bold.ttf")
        if bold_path.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(bold_path)))
        logger.info("Registered DejaVu for ReportLab: %s", DEJAVU_PATH)
    except Exception as e:
        logger.exception("Failed to register DejaVu for ReportLab: %s", e)
else:
    logger.info("No DejaVu font found in %s", str(FONTS_DIR))


# ---------------- IMAGE SCALING UTIL ---------------------------------------
def scaled_rl_image(path: Path, max_w_mm: float = 36.0, max_h_mm: float = 36.0):
    """
    Return reportlab Image flowable scaled to max width/height (mm) preserving aspect ratio.
    """
    if not Path(path).exists():
        raise FileNotFoundError(f"{path} not found")
    try:
        reader = ImageReader(str(path))
        px_w, px_h = reader.getSize()  # pixel dims
        # target points
        max_w_pts = max_w_mm * mm
        max_h_pts = max_h_mm * mm
        # scale factor to fit within box
        if px_w <= 0 or px_h <= 0:
            scale = 1.0
        else:
            # treat px as points baseline; compute scale
            scale = min(max_w_pts / px_w, max_h_pts / px_h, 1.0)
        img = RLImage(str(path))
        img.drawWidth = px_w * scale
        img.drawHeight = px_h * scale
        return img
    except Exception:
        # fallback: basic image object without scaling
        img = RLImage(str(path))
        # limit hard-coded small fallback
        img.drawWidth = max_w_mm * mm
        img.drawHeight = max_h_mm * mm
        return img


# ---------------- PDF GENERATOR (Platypus + fallback) -----------------------
def _wrap_text_simple(text, chars_per_line=90):
    words = str(text).split()
    lines = []
    cur = ""
    for w in words:
        if len(cur) + len(w) + 1 <= chars_per_line:
            cur = cur + (" " if cur else "") + w
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def _fallback_canvas_pdf(
    patient,
    prakriti_pct,
    vikriti_pct,
    psych_pct,
    career_recs,
    rel_tips,
    health_recs,
    error_text=None,
    report_id=None,
):
    """
    Readable canvas fallback (BytesIO). Includes error snippet.
    """
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    W, H = A4
    left = 24 * mm
    top = H - 24 * mm
    try:
        if DEJAVU_PATH:
            c.setFont("DejaVuSans", 10)
        else:
            c.setFont("Helvetica", 10)
    except Exception:
        c.setFont("Helvetica", 10)
    y = top

    logo_path = APP_DIR / "logo.png"
    if not logo_path.exists() and Path("logo.png").exists():
        logo_path = Path("logo.png")
    if logo_path.exists():
        try:
            reader = ImageReader(str(logo_path))
            px_w, px_h = reader.getSize()
            scale = min((36 * mm) / px_w, (36 * mm) / px_h, 1.0)
            c.drawImage(
                str(logo_path),
                left,
                y - (36 * mm),
                width=px_w * scale,
                height=px_h * scale,
                preserveAspectRatio=True,
                mask="auto",
            )
        except Exception:
            pass

    c.setFont("Helvetica-Bold", 13)
    c.drawString(left + 140, y - 6, BRAND.get("clinic_name", ""))
    c.setFont("Helvetica", 9)
    c.drawString(left + 140, y - 20, BRAND.get("tagline", ""))
    y -= 46 * mm

    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y, f"Patient: {patient.get('name','')}")
    c.setFont("Helvetica", 9)
    c.drawString(left + 220, y, f"Date: {datetime.now().strftime('%Y-%m-%d')}")
    y -= 12

    summary = "Fallback report: primary generator failed. Below are the generated recommendations and data (layout simplified)."
    for line in _wrap_text_simple(summary, 95):
        c.drawString(left, y, line)
        y -= 10

    y -= 6
    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y, "Prakriti:")
    y -= 10
    c.setFont("Helvetica", 9)
    for k, v in prakriti_pct.items():
        c.drawString(left + 6, y, f"{k}: {v} %")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top

    y -= 6
    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y, "Vikriti:")
    y -= 10
    c.setFont("Helvetica", 9)
    for k, v in vikriti_pct.items():
        c.drawString(left + 6, y, f"{k}: {v} %")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top

    y -= 6
    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y, "Top career suggestions:")
    y -= 10
    c.setFont("Helvetica", 9)
    for cr in career_recs[:10]:
        c.drawString(left + 6, y, f"- {cr.get('role')} (score {cr.get('score')})")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top

    y -= 6
    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y, "Relationship tips:")
    y -= 10
    c.setFont("Helvetica", 9)
    for t in rel_tips:
        c.drawString(left + 6, y, f"- {t[0]}: {t[1]}")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top

    y -= 6
    c.setFont("Helvetica-Bold", 10)
    c.drawString(left, y, "Health suggestions:")
    y -= 10
    c.setFont("Helvetica", 9)
    for d in health_recs.get("diet", []):
        c.drawString(left + 6, y, f"- Diet: {d}")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top
    for l in health_recs.get("lifestyle", []):
        c.drawString(left + 6, y, f"- Lifestyle: {l}")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top
    for h in health_recs.get("herbs", []):
        c.drawString(left + 6, y, f"- Herbs: {h}")
        y -= 10
        if y < 70 * mm:
            c.showPage()
            y = top

    if error_text:
        y -= 8
        if y < 80 * mm:
            c.showPage()
            y = top
        c.setFont("Helvetica-Bold", 9)
        c.drawString(left, y, "Report engine error (short):")
        y -= 12
        c.setFont("Helvetica", 7)
        for line in _wrap_text_simple(error_text, 120):
            c.drawString(left, y, line)
            y -= 8
            if y < 30 * mm:
                c.showPage()
                y = top

    c.setFont("Helvetica", 8)
    c.drawString(
        left,
        18 * mm,
        f"{BRAND.get('clinic_name')} — {BRAND.get('phone')} — {BRAND.get('email')}",
    )
    c.save()
    buf.seek(0)
    return buf


def branded_pdf_report(
    patient,
    prakriti_pct,
    vikriti_pct,
    psych_pct,
    career_recs,
    rel_tips,
    health_recs,
    report_id=None,
):
    """
    Main branded PDF builder. Returns BytesIO.
    Tries Platypus flow; on failure returns canvas-based fallback BytesIO.
    """
    logger.info("Start branded_pdf_report for %s", patient.get("name", "N/A"))
    # create chart images
    p1 = TMP_DIR / f"prakriti_{int(datetime.now().timestamp())}.png"
    p2 = TMP_DIR / f"vikriti_{int(datetime.now().timestamp())}.png"
    p3 = TMP_DIR / f"psych_{int(datetime.now().timestamp())}.png"
    try:
        _make_bar_chart(prakriti_pct, "Prakriti (constitutional %)", p1)
        _make_bar_chart(vikriti_pct, "Vikriti (today %)", p2)
        _make_bar_chart(psych_pct, "Psychometric (approx %)", p3)
    except Exception:
        logger.exception("Chart generation failed")

    try:
        buf = BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
        )
        styles = getSampleStyleSheet()
        base_font = "DejaVuSans" if DEJAVU_PATH else "Helvetica"
        accent = (
            colors.HexColor(BRAND["accent_color"])
            if BRAND.get("accent_color")
            else colors.HexColor("#0F7A61")
        )

        # add unique styles if not present
        if "AP_Title" not in styles:
            styles.add(
                ParagraphStyle(
                    name="AP_Title",
                    fontName=base_font if base_font else "Helvetica",
                    fontSize=18,
                    leading=22,
                    spaceAfter=6,
                )
            )
        if "AP_Small" not in styles:
            styles.add(
                ParagraphStyle(
                    name="AP_Small",
                    fontName=base_font if base_font else "Helvetica",
                    fontSize=9,
                    leading=11,
                )
            )
        if "AP_Heading" not in styles:
            styles.add(
                ParagraphStyle(
                    name="AP_Heading",
                    fontName=base_font if base_font else "Helvetica",
                    fontSize=12,
                    leading=14,
                    spaceBefore=8,
                    spaceAfter=4,
                    textColor=accent,
                )
            )
        if "AP_Body" not in styles:
            styles.add(
                ParagraphStyle(
                    name="AP_Body",
                    fontName=base_font if base_font else "Helvetica",
                    fontSize=10,
                    leading=13,
                )
            )
        if "AP_Bullet" not in styles:
            styles.add(
                ParagraphStyle(
                    name="AP_Bullet",
                    fontName=base_font if base_font else "Helvetica",
                    fontSize=10,
                    leading=12,
                    leftIndent=12,
                    bulletIndent=6,
                )
            )

        flow = []

        # Header: logo + clinic info
        logo_path = APP_DIR / "logo.png"
        if not logo_path.exists() and Path("logo.png").exists():
            logo_path = Path("logo.png")
        if logo_path.exists():
            try:
                img = scaled_rl_image(logo_path, max_w_mm=36, max_h_mm=36)
                clinic_info = Paragraph(
                    f"<b>{BRAND['clinic_name']}</b><br/>{BRAND['tagline']}<br/><font size=9>{BRAND['website']}</font>",
                    styles["AP_Body"],
                )
                header_t = Table([[img, clinic_info]], colWidths=[40 * mm, 120 * mm])
                header_t.setStyle(
                    TableStyle(
                        [
                            ("VALIGN", (0, 0), (1, 0), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ]
                    )
                )
                flow.append(header_t)
            except Exception:
                logger.exception("Logo embedding failed in Platypus")
                flow.append(
                    Paragraph(
                        f"<b>{BRAND['clinic_name']}</b><br/>{BRAND['tagline']}",
                        styles["AP_Title"],
                    )
                )
        else:
            flow.append(
                Paragraph(
                    f"<b>{BRAND['clinic_name']}</b><br/>{BRAND['tagline']}",
                    styles["AP_Title"],
                )
            )

        flow.append(Spacer(1, 6))

        # Patient details table
        patient_lines = [
            ["Patient", patient.get("name", "")],
            ["Age", str(patient.get("age", ""))],
            ["Gender", patient.get("gender", "")],
            ["Date", datetime.now().strftime("%Y-%m-%d")],
            ["Report ID", str(report_id or "N/A")],
        ]
        t = Table(patient_lines, colWidths=[70 * mm, 80 * mm])
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (1, 0), colors.whitesmoke),
                    ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ("BOX", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        flow.append(t)
        flow.append(Spacer(1, 8))

        # Executive summary
        flow.append(Paragraph("Executive summary", styles["AP_Heading"]))
        summary_text = (
            "This report summarises constitutional profile (Prakriti), current imbalances (Vikriti) and "
            "a psychometric snapshot. It includes practical, prioritized recommendations for career, "
            "relationships and health."
        )
        flow.append(Paragraph(summary_text, styles["AP_Body"]))
        flow.append(Spacer(1, 8))

        # Charts (two across, one below)
        try:
            cells = []
            if p1.exists() and p2.exists():
                img1 = RLImage(str(p1), width=85 * mm, height=45 * mm)
                img2 = RLImage(str(p2), width=85 * mm, height=45 * mm)
                cells.append([img1, img2])
            if cells:
                flow.append(
                    Table(
                        cells,
                        colWidths=[90 * mm, 90 * mm],
                        style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE")],
                    )
                )
                flow.append(Spacer(1, 6))
            if p3.exists():
                img3 = RLImage(str(p3), width=100 * mm, height=45 * mm)
                flow.append(img3)
                flow.append(Spacer(1, 6))
        except Exception:
            logger.exception("Failed to add chart images to the Platypus flow")

        # Prakriti & Vikriti numeric tables
        flow.append(
            Paragraph("Prakriti — percentage distribution", styles["AP_Heading"])
        )
        pp = [[k, f"{v} %"] for k, v in prakriti_pct.items()]
        t = Table(pp, colWidths=[60 * mm, 40 * mm])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        flow.append(t)
        flow.append(Spacer(1, 6))

        flow.append(
            Paragraph("Vikriti — percentage distribution (today)", styles["AP_Heading"])
        )
        vp = [[k, f"{v} %"] for k, v in vikriti_pct.items()]
        t = Table(vp, colWidths=[60 * mm, 40 * mm])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        flow.append(t)
        flow.append(Spacer(1, 8))

        # Recommendations block
        flow.append(
            Paragraph("Recommendations (practical & prioritized)", styles["AP_Heading"])
        )
        flow.append(Paragraph("Career suggestions (ranked):", styles["AP_Body"]))
        for cr in career_recs:
            role = cr.get("role", "-")
            score = cr.get("score", 0)
            reason = cr.get("reason", "")
            flow.append(
                Paragraph(
                    f"• <b>{role}</b> — score: {score}. <i>{reason}</i>",
                    styles["AP_Bullet"],
                )
            )
        flow.append(Spacer(1, 6))

        flow.append(Paragraph("Relationship & communication tips:", styles["AP_Body"]))
        for title, reason in rel_tips:
            flow.append(Paragraph(f"• <b>{title}</b> — {reason}", styles["AP_Bullet"]))
        flow.append(Spacer(1, 6))

        flow.append(
            Paragraph("Health, Diet & Lifestyle (practical):", styles["AP_Body"])
        )
        if health_recs:
            flow.append(Paragraph("<b>Diet</b>:", styles["AP_Small"]))
            for d in health_recs.get("diet", []):
                flow.append(Paragraph(f"• {d}", styles["AP_Bullet"]))
            flow.append(Paragraph("<b>Lifestyle</b>:", styles["AP_Small"]))
            for l in health_recs.get("lifestyle", []):
                flow.append(Paragraph(f"• {l}", styles["AP_Bullet"]))
            flow.append(
                Paragraph("<b>Herbs (consult physician)</b>:", styles["AP_Small"])
            )
            for h in health_recs.get("herbs", []):
                flow.append(Paragraph(f"• {h}", styles["AP_Bullet"]))
        flow.append(Spacer(1, 8))

        # Footer contact
        flow.append(Spacer(1, 12))
        contact_par = f"{BRAND.get('clinic_name')} — {BRAND.get('doctor')} — {BRAND.get('phone')} — {BRAND.get('email')}"
        flow.append(Paragraph(contact_par, styles["AP_Small"]))
        flow.append(Paragraph(BRAND.get("address", ""), styles["AP_Small"]))

        doc.build(flow)
        buf.seek(0)

        # cleanup charts
        for p in (p1, p2, p3):
            try:
                if p.exists():
                    p.unlink()
            except Exception:
                pass

        logger.info(
            "Platypus PDF built successfully for %s", patient.get("name", "N/A")
        )
        return buf

    except Exception:
        tb = traceback.format_exc()
        logger.exception("Platypus build failed: %s", tb)
        snippet = tb[:1200]
        return _fallback_canvas_pdf(
            patient,
            prakriti_pct,
            vikriti_pct,
            psych_pct,
            career_recs,
            rel_tips,
            health_recs,
            error_text=snippet,
            report_id=report_id,
        )


# ---------------- DOCX REPORT (simple) -------------------------------------
def docx_report(
    patient, prakriti_pct, vikriti_pct, psych_pct, career_recs, rel_tips, health_recs
):
    doc = Document()
    doc.add_heading(f"{CONFIG['meta']['app_name']} — Personalized Report", level=1)
    doc.add_paragraph(
        f"Name: {patient.get('name')}    Age: {patient.get('age')}    Gender: {patient.get('gender')}"
    )
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_heading("Prakriti (constitutional) %", level=2)
    for k, v in prakriti_pct.items():
        doc.add_paragraph(f"{k}: {v} %", style="List Bullet")
    doc.add_heading("Vikriti (today) %", level=2)
    for k, v in vikriti_pct.items():
        doc.add_paragraph(f"{k}: {v} %", style="List Bullet")
    doc.add_heading("Psychometric summary (approx)", level=2)
    for k, v in psych_pct.items():
        doc.add_paragraph(f"{k}: {v} %", style="List Bullet")
    doc.add_heading("Recommendations", level=2)
    doc.add_heading("Career suggestions (ranked)", level=3)
    for cr in career_recs:
        doc.add_paragraph(f"{cr['role']} (score: {cr['score']})", style="List Number")
    doc.add_heading("Relationship tips", level=3)
    for t in rel_tips:
        doc.add_paragraph(t[0] + " — " + t[1], style="List Bullet")
    doc.add_heading("Health & lifestyle", level=3)
    doc.add_paragraph("Diet: " + ", ".join(health_recs["diet"]))
    doc.add_paragraph("Lifestyle: " + ", ".join(health_recs["lifestyle"]))
    doc.add_paragraph("Herbs: " + ", ".join(health_recs["herbs"]))
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# ---------------- STREAMLIT UI ---------------------------------------------
st.set_page_config(page_title=CONFIG["meta"]["app_name"], layout="wide")
st.markdown(
    "<style>section[data-testid='stSidebar'] {background-color: #f7f7fa}</style>",
    unsafe_allow_html=True,
)

col1, col2 = st.columns([6, 4])
with col1:
    st.title(CONFIG["meta"]["app_name"])
    st.caption(f"Version {CONFIG['meta']['version']} — {CONFIG['meta']['author']}")
with col2:
    st.markdown("**Session**")
    st.write(datetime.now().strftime("%Y-%m-%d %H:%M"))

# Authentication sidebar
st.sidebar.subheader("Login")
username = st.sidebar.text_input("Username")
password = st.sidebar.text_input("Password", type="password")
if "auth" not in st.session_state:
    st.session_state.auth = False
if st.sidebar.button("Login"):
    ok, info = verify_user(username, password)
    if ok:
        st.session_state.auth = True
        st.session_state.user = username
        st.session_state.user_info = info
        st.sidebar.success(f"Welcome {info['display_name']} ({info['role']})")
    else:
        st.sidebar.error("Invalid username/password")
if not st.session_state.auth:
    st.info(
        "Please login from the sidebar. Use default admin / admin123 for first login (change it later)."
    )
    st.stop()
st.sidebar.markdown("---")
role = st.session_state.user_info.get("role", "clinician")
st.sidebar.write("Role: " + role)

tabs = st.tabs(
    ["Patient Registry", "New Assessment", "Clinician Dashboard", "Config & Export"]
)

# Tab 1: Patient registry
with tabs[0]:
    st.header("Patient registry")
    with st.expander("Create new patient"):
        with st.form("new_patient"):
            pname = st.text_input("Full name")
            page = st.number_input("Age", min_value=0, max_value=120, value=30)
            pgender = st.selectbox(
                "Gender", ["Male", "Female", "Other", "Prefer not to say"]
            )
            pcontact = st.text_input("Contact (phone/email)")
            if st.form_submit_button("Create patient"):
                if not pname:
                    st.warning("Enter a name")
                else:
                    pid = create_patient(pname, page, pgender, pcontact)
                    st.success(f"Patient created (id: {pid})")
    st.write("### Existing patients")
    patients_df = load_patients()
    st.dataframe(patients_df)

# Tab 2: New Assessment
with tabs[1]:
    st.header("New Assessment — Prakriti, Vikriti & Psychometrics")
    patients = load_patients()
    if patients.empty:
        st.info("No patients found. Create a patient first.")
    else:
        psel = st.selectbox(
            "Select patient",
            options=patients["id"].tolist(),
            format_func=lambda x: f"{int(x)} - {patients.loc[patients['id']==x,'name'].values[0]}",
        )
        patient_row = patients[patients["id"] == psel].iloc[0].to_dict()
        st.markdown(
            f"**Patient:** {patient_row['name']} | Age: {patient_row['age']} | Gender: {patient_row['gender']}"
        )
        st.markdown("---")
        pr_qs = CONFIG["questions"]["prakriti"]
        pr_answers = {}
        cols = st.columns(2)
        for i, q in enumerate(pr_qs):
            with cols[i % 2]:
                pr_answers[q["id"]] = st.slider(q["text"], 1, 5, 3, key=f"pr_{q['id']}")
        vk_qs = CONFIG["questions"]["vikriti"]
        vk_answers = {}
        cols = st.columns(3)
        for i, q in enumerate(vk_qs):
            with cols[i % 3]:
                vk_answers[q["id"]] = st.slider(q["text"], 1, 5, 1, key=f"vk_{q['id']}")
        psy_qs = CONFIG["questions"]["psychometric"]
        psy_answers = {}
        cols = st.columns(2)
        for i, q in enumerate(psy_qs):
            with cols[i % 2]:
                psy_answers[q["id"]] = st.slider(
                    q["text"], 1, 7, 4, key=f"psy_{q['id']}"
                )
        if st.button("Compute & Save Assessment"):
            prak_pct = score_dosha_from_answers(pr_answers, pr_qs)
            vik_pct = score_dosha_from_answers(vk_answers, vk_qs)
            psych_pct = psychometric_tipiscale(psy_answers)
            career = recommend_career(prak_pct, psych_pct)
            rel = recommend_relationship(prak_pct, psych_pct)
            health = recommend_health(prak_pct, vik_pct)
            payload = {
                "patient": patient_row,
                "prakriti_answers": pr_answers,
                "vikriti_answers": vk_answers,
                "psych_answers": psy_answers,
                "prakriti_pct": prak_pct,
                "vikriti_pct": vik_pct,
                "psych_pct": psych_pct,
                "career_recs": career,
                "relationship_tips": rel,
                "health_recs": health,
            }
            aid = save_assessment(patient_row["id"], st.session_state.user, payload)
            st.success(f"Assessment saved (id: {aid})")
            st.markdown("### Results snapshot")
            c1, c2, c3 = st.columns(3)
            c1.metric("Dominant Prakriti", max(prak_pct, key=prak_pct.get))
            c2.metric("Current Aggravation", max(vik_pct, key=vik_pct.get))
            c3.metric("Dominant Trait", max(psych_pct, key=psych_pct.get))
            st.write("#### Prakriti (percent)")
            st.bar_chart(pd.Series(prak_pct))
            st.write("#### Vikriti (percent)")
            st.bar_chart(pd.Series(vik_pct))
            st.write("#### Psychometric")
            st.bar_chart(pd.Series(psych_pct))
            st.write("#### Career recommendations (top 5)")
            for r in career[:5]:
                st.write(
                    f"- {r['role']} (score: {r['score']})  —  {r.get('reason','')}"
                )
            st.write("#### Relationship tips")
            for t in rel:
                st.write("- " + t[0] + " — " + t[1])
            st.write("#### Health suggestions")
            st.write("Diet: " + ", ".join(health["diet"]))
            st.write("Lifestyle: " + ", ".join(health["lifestyle"]))
            st.write("Herbs: " + ", ".join(health["herbs"]))

            # PDF (BytesIO) and DOCX
            pdf_b = branded_pdf_report(
                patient_row,
                prak_pct,
                vik_pct,
                psych_pct,
                career,
                rel,
                health,
                report_id=aid,
            )
            st.download_button(
                "Download Branded PDF (professional)",
                pdf_b,
                file_name=f"Branded_Report_{patient_row['name']}_{aid}.pdf",
                mime="application/pdf",
            )
            docx_b = docx_report(
                patient_row, prak_pct, vik_pct, psych_pct, career, rel, health
            )
            st.download_button(
                "Download DOCX report",
                docx_b,
                file_name=f"Report_{patient_row['name']}_{aid}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )

            follow_date = (datetime.now() + timedelta(days=7)).strftime("%Y%m%dT%H%M00")
            ics = f"BEGIN:VCALENDAR\nVERSION:2.0\nBEGIN:VEVENT\nDTSTAMP:{datetime.now().strftime('%Y%m%dT%H%M00')}\nDTSTART:{follow_date}\nDTEND:{follow_date}\nSUMMARY:Follow-up for {patient_row['name']} (AyurPrakriti)\nDESCRIPTION:Review outcomes and adherence to recommendations.\nEND:VEVENT\nEND:VCALENDAR"
            st.download_button(
                "Download follow-up (.ics)",
                ics.encode("utf-8"),
                file_name=f"followup_{patient_row['name']}_{aid}.ics",
                mime="text/calendar",
            )

# Tab 3: Clinician Dashboard
with tabs[2]:
    st.header("Clinician Dashboard")
    st.write("Recent assessments")
    asses = load_assessments()
    if asses.empty:
        st.info("No assessments yet")
    else:
        st.dataframe(asses[["id", "patient_id", "assessor", "created_at"]].head(50))
        sel = st.number_input("Open assessment id", min_value=0, value=0, step=1)
        if sel > 0:
            cur.execute("SELECT data_json FROM assessments WHERE id=?", (int(sel),))
            r = cur.fetchone()
            if r:
                st.json(json.loads(r[0]))
            else:
                st.warning("Not found")
    st.markdown("---")
    st.subheader("Manage users (admin only)")
    if role != "admin":
        st.info("User management visible to admin users only")
    else:
        with st.form("create_user"):
            un = st.text_input("Username")
            dn = st.text_input("Display name")
            pw = st.text_input("Password", type="password")
            rrole = st.selectbox("Role", ["clinician", "admin"])
            if st.form_submit_button("Create user"):
                if not un or not pw:
                    st.warning("Provide username & password")
                else:
                    ph = pwd_context.hash(pw)
                    try:
                        cur.execute(
                            "INSERT INTO users (username, display_name, password_hash, role, created_at) VALUES (?,?,?,?,?)",
                            (un, dn, ph, rrole, datetime.now().isoformat()),
                        )
                        conn.commit()
                        st.success("User created")
                    except Exception as e:
                        st.error("Error: " + str(e))

# Tab 4: Config & Export
with tabs[3]:
    st.header("Config & Export")
    st.subheader("Branding")
    st.write(
        "Upload a logo (PNG/JPG) to include in the branded PDF. Current logo will be detected from app folder if present."
    )
    logo_file = st.file_uploader("Upload logo.png", type=["png", "jpg", "jpeg"])
    if logo_file is not None:
        save_path = APP_DIR / "logo.png"
        with open(save_path, "wb") as f:
            f.write(logo_file.getbuffer())
        st.success("Logo uploaded and saved (will appear in next report)")
    st.subheader("Clinic details (branding)")
    with st.form("branding"):
        BRAND["clinic_name"] = st.text_input(
            "Clinic / Brand Name", value=BRAND["clinic_name"]
        )
        BRAND["tagline"] = st.text_input("Tagline", value=BRAND["tagline"])
        BRAND["doctor"] = st.text_input(
            "Doctor Name & Qualifications", value=BRAND["doctor"]
        )
        BRAND["address"] = st.text_input("Clinic Address", value=BRAND["address"])
        BRAND["phone"] = st.text_input("Phone", value=BRAND["phone"])
        BRAND["email"] = st.text_input("Email", value=BRAND["email"])
        BRAND["website"] = st.text_input("Website", value=BRAND["website"])
        BRAND["accent_color"] = st.text_input(
            "Accent color (hex)", value=BRAND["accent_color"]
        )
        if st.form_submit_button("Save branding"):
            st.success("Branding updated")
    st.markdown("---")
    st.subheader("Config file (editable YAML)")
    cfg_text = yaml.safe_dump(CONFIG, sort_keys=False)
    new_cfg_text = st.text_area("Edit YAML config", cfg_text, height=300)
    if st.button("Save config"):
        try:
            newcfg = yaml.safe_load(new_cfg_text)
            with open(CFG_PATH, "w", encoding="utf-8") as f:
                yaml.safe_dump(newcfg, f, sort_keys=False)
            st.success("Config saved. Please restart the app for full effect.")
        except Exception as e:
            st.error("Invalid YAML: " + str(e))
    st.markdown("---")
    st.subheader("Export DB (download)")
    conn.commit()
    with open(DB_PATH, "rb") as f:
        dbdata = f.read()
    st.download_button(
        "Download SQLite DB",
        data=dbdata,
        file_name="ayurprakriti.db",
        mime="application/octet-stream",
    )
    st.markdown("---")
    st.caption(
        "Next steps: migrate DB to Postgres for multi-user, add OAuth, add consent & audit logs, host behind HTTPS."
    )

# ---------------- END OF FILE ------------------------------------------------
